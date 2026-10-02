from flask import Flask, render_template, request, jsonify, send_file
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from copy import copy
from datetime import datetime
import os
import io
from reconciliation import reconcile_data, format_mst

app = Flask(__name__)
UPLOAD_FOLDER = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'uploads')
os.makedirs(UPLOAD_FOLDER, exist_ok=True)

def create_4sheet_export_workbook(df_orders, df_items, df_unmatched_inv, invoice_row_matches, invoice_path):
    """
    Build a 4-sheet Excel workbook:
    - Sheet 1: Tổng quan đơn hàng
    - Sheet 2: Chi tiết đối chiếu
    - Sheet 3: Hóa đơn chưa khớp Sapo
    - Sheet 4: Sheet hóa đơn gốc (NK_HDXML) giữ nguyên 100% mẫu biểu của Tổng cục Thuế,
               tự động điền mã SKU Sapo vào cột 'Mã hàng' và thêm cột 'Mã đơn nhập hàng' sau cột 'Ghi chú'.
    """
    # 1. Load original invoice workbook to preserve 100% template, fonts, row heights, and styles
    wb = openpyxl.load_workbook(invoice_path)
    ws_inv = wb.active
    
    # Locate header row containing "Số hóa đơn"
    header_row_idx = None
    for r in range(1, 15):
        for c in range(1, ws_inv.max_column + 1):
            if str(ws_inv.cell(row=r, column=c).value).strip() == "Số hóa đơn":
                header_row_idx = r
                break
        if header_row_idx:
            break
            
    col_so_hd = None
    col_ma_hang = None
    col_ghi_chu = None
    for c in range(1, ws_inv.max_column + 1):
        val = str(ws_inv.cell(row=header_row_idx, column=c).value).strip()
        if val == "Số hóa đơn":
            col_so_hd = c
        elif val == "Mã hàng":
            col_ma_hang = c
        elif val == "Ghi chú":
            col_ghi_chu = c
            
    # Target column for new column "Mã đơn nhập hàng" (immediately after Ghi chú)
    col_ma_don = (col_ghi_chu + 1) if col_ghi_chu else (ws_inv.max_column + 1)
    
    # Unmerge any merged cells overlapping on col_ma_don around header rows (e.g. Q5:Q6)
    for m in list(ws_inv.merged_cells.ranges):
        min_col, min_row, max_col, max_row = m.bounds
        if min_col <= col_ma_don <= max_col and (min_row <= header_row_idx <= max_row or min_row <= header_row_idx - 1 <= max_row):
            ws_inv.unmerge_cells(m.coord)
            
    # Add new column headers: "Mã đơn nhập hàng" and "[16]"
    cell_h1 = ws_inv.cell(row=header_row_idx, column=col_ma_don, value="Mã đơn nhập hàng")
    if col_ghi_chu:
        ref_h1 = ws_inv.cell(row=header_row_idx, column=col_ghi_chu)
        if ref_h1.has_style:
            cell_h1.font = copy(ref_h1.font)
            cell_h1.fill = copy(ref_h1.fill)
            cell_h1.alignment = copy(ref_h1.alignment)
            cell_h1.border = copy(ref_h1.border)
            
    cell_h2 = ws_inv.cell(row=header_row_idx + 1, column=col_ma_don, value="[16]")
    if col_ghi_chu:
        ref_h2 = ws_inv.cell(row=header_row_idx + 1, column=col_ghi_chu)
        if ref_h2.has_style:
            cell_h2.font = copy(ref_h2.font)
            cell_h2.fill = copy(ref_h2.fill)
            cell_h2.alignment = copy(ref_h2.alignment)
            cell_h2.border = copy(ref_h2.border)
            
    ws_inv.column_dimensions[get_column_letter(col_ma_don)].width = 22
    
    # Populate matched SKUs into 'Mã hàng' and Order IDs into 'Mã đơn nhập hàng'
    for idx, match_info in invoice_row_matches.items():
        target_row = header_row_idx + 1 + idx
        if col_ma_hang:
            ws_inv.cell(row=target_row, column=col_ma_hang, value=match_info["sku"])
        cell_don = ws_inv.cell(row=target_row, column=col_ma_don, value=match_info["order_id"])
        if col_ghi_chu:
            ref_cell = ws_inv.cell(row=target_row, column=col_ghi_chu)
            if ref_cell.has_style:
                cell_don.font = copy(ref_cell.font)
                cell_don.alignment = Alignment(horizontal="center", vertical="center")
                cell_don.border = copy(ref_cell.border)
                
    # 2. Insert summary sheets at indices 0, 1, 2
    ws_orders = wb.create_sheet(title='Tổng quan đơn hàng', index=0)
    ws_items = wb.create_sheet(title='Chi tiết đối chiếu', index=1)
    ws_unmatched = wb.create_sheet(title='Hóa đơn chưa khớp Sapo', index=2)
    
    # Styles for summary sheets
    font_header = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    fill_header = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_right = Alignment(horizontal="right", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )
    fill_green = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    font_green = Font(name="Arial", size=10, color="006100")
    fill_yellow = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
    font_yellow = Font(name="Arial", size=10, color="9C6500")
    fill_red = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    font_red = Font(name="Arial", size=10, color="9C0006")
    font_regular = Font(name="Arial", size=10)
    
    def populate_sheet(ws, df, col_names, money_cols, status_col=None):
        ws.row_dimensions[1].height = 26
        for col_idx, col_name in enumerate(col_names, 1):
            cell = ws.cell(row=1, column=col_idx, value=col_name)
            cell.font = font_header
            cell.fill = fill_header
            cell.alignment = align_center
            cell.border = thin_border
            
        for row_idx, row_data in enumerate(df.values, 2):
            ws.row_dimensions[row_idx].height = 20
            for col_idx, val in enumerate(row_data, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=val)
                cell.font = font_regular
                cell.border = thin_border
                
                if col_idx in money_cols:
                    try:
                        cell.value = float(val)
                        cell.number_format = '#,##0'
                    except:
                        pass
                    cell.alignment = align_right
                else:
                    cell.alignment = align_left
                    
                if status_col and col_idx == status_col:
                    s_val = str(val).strip()
                    if "Khớp" in s_val:
                        cell.fill = fill_green
                        cell.font = font_green
                    elif "Lệch" in s_val:
                        cell.fill = fill_yellow
                        cell.font = font_yellow
                    elif "Chưa" in s_val or "Thiếu" in s_val:
                        cell.fill = fill_red
                        cell.font = font_red
                        
        for col in ws.columns:
            col_letter = get_column_letter(col[0].column)
            max_len = max(len(str(c.value or '')) for c in col)
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 45)
            
    # Sheet 1: Tổng quan đơn hàng
    df_orders_excel = df_orders[[
        'order_id', 'date', 'supplier_code', 'supplier_name', 'mst', 'total_incl', 'total_excl', 'status', 'notes'
    ]]
    order_headers = [
        'Mã đơn nhập hàng', 'Ngày nhập', 'Mã NCC (Sapo)', 'Tên NCC (Sapo)', 'Mã số thuế',
        'Tổng tiền (Có VAT)', 'Tổng tiền (Không VAT)', 'Trạng thái đối chiếu', 'Chi tiết sai lệch'
    ]
    populate_sheet(ws_orders, df_orders_excel, order_headers, money_cols=[6, 7], status_col=8)
    
    # Sheet 2: Chi tiết đối chiếu
    df_items_excel = df_items[[
        'order_id', 'sku', 'name', 'qty', 'price_excl', 'total_excl',
        'invoice_no', 'inv_name', 'inv_qty', 'inv_price', 'inv_total', 'status', 'diff_qty', 'diff_amt'
    ]]
    item_headers = [
        'Mã đơn nhập hàng', 'Mã SKU Sapo', 'Tên sản phẩm Sapo', 'SL nhập (Sapo)',
        'Đơn giá chưa thuế (Sapo)', 'Thành tiền chưa thuế (Sapo)', 'Số hóa đơn TCT',
        'Tên hàng trên hóa đơn', 'SL hóa đơn (TCT)', 'Đơn giá hóa đơn (TCT)',
        'Thành tiền hóa đơn (TCT)', 'Kết quả', 'Lệch số lượng', 'Lệch thành tiền'
    ]
    populate_sheet(ws_items, df_items_excel, item_headers, money_cols=[5, 6, 10, 11, 14], status_col=12)
    
    # Sheet 3: Hóa đơn chưa khớp Sapo
    df_unmatched_excel = df_unmatched_inv[[
        'mst', 'supplier_name', 'date', 'invoice_no', 'name', 'qty', 'price', 'total'
    ]]
    unmatched_headers = [
        'Mã số thuế người bán', 'Tên nhà cung cấp', 'Ngày hóa đơn', 'Số hóa đơn TCT',
        'Tên hàng hóa TCT', 'Số lượng', 'Đơn giá chưa thuế', 'Thành tiền chưa thuế'
    ]
    populate_sheet(ws_unmatched, df_unmatched_excel, unmatched_headers, money_cols=[7, 8])
    
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output

@app.route('/')
def index():
    return render_template('index.html')

@app.route('/reconcile', methods=['POST'])
def handle_reconcile():
    if 'sapo_file' not in request.files or 'supplier_file' not in request.files or 'invoice_file' not in request.files:
        return jsonify({'error': 'Vui lòng tải lên đầy đủ 3 file nguồn.'}), 400
        
    sapo_file = request.files['sapo_file']
    supplier_file = request.files['supplier_file']
    invoice_file = request.files['invoice_file']
    
    if sapo_file.filename == '' or supplier_file.filename == '' or invoice_file.filename == '':
        return jsonify({'error': 'Một hoặc nhiều file chưa được chọn.'}), 400
        
    try:
        # Read bytes
        sapo_bytes = sapo_file.read()
        supplier_bytes = supplier_file.read()
        invoice_bytes = invoice_file.read()
        
        # Save files temporarily for potential export rebuild
        sapo_path = os.path.join(UPLOAD_FOLDER, 'sapo_latest.xlsx')
        supplier_path = os.path.join(UPLOAD_FOLDER, 'supplier_latest.xlsx')
        invoice_path = os.path.join(UPLOAD_FOLDER, 'invoice_latest.xlsx')
        
        with open(sapo_path, 'wb') as f:
            f.write(sapo_bytes)
        with open(supplier_path, 'wb') as f:
            f.write(supplier_bytes)
        with open(invoice_path, 'wb') as f:
            f.write(invoice_bytes)
            
        # Reconcile
        df_orders, df_items, df_unmatched_inv, _ = reconcile_data(sapo_path, supplier_path, invoice_path)
        
        # Check matching summary stats
        total_orders = len(df_orders)
        stats = {
            'total': total_orders,
            'khop': int((df_orders['status'] == 'Khớp').sum()),
            'lech': int((df_orders['status'] == 'Lệch').sum()),
            'chua_hd': int((df_orders['status'] == 'Chưa xuất hóa đơn').sum()),
            'thieu_mst': int((df_orders['status'] == 'Thiếu MST').sum())
        }
        
        # Convert DataFrames to dict lists
        orders_list = df_orders.to_dict(orient='records')
        items_list = df_items.to_dict(orient='records')
        unmatched_list = df_unmatched_inv.to_dict(orient='records')
        
        return jsonify({
            'success': True,
            'stats': stats,
            'orders': orders_list,
            'items': items_list,
            'unmatched_invoices': unmatched_list
        })
        
    except Exception as e:
        import traceback
        traceback.print_exc()
        return jsonify({'error': f'Đã xảy ra lỗi khi đối chiếu: {str(e)}'}), 500

@app.route('/export', methods=['POST'])
def handle_export():
    try:
        # Load from temp saved paths to rebuild dataframes and export beautifully
        sapo_path = os.path.join(UPLOAD_FOLDER, 'sapo_latest.xlsx')
        supplier_path = os.path.join(UPLOAD_FOLDER, 'supplier_latest.xlsx')
        invoice_path = os.path.join(UPLOAD_FOLDER, 'invoice_latest.xlsx')
        
        if not (os.path.exists(sapo_path) and os.path.exists(supplier_path) and os.path.exists(invoice_path)):
            return "Không tìm thấy dữ liệu đối chiếu trước đó. Vui lòng tải lại file.", 400
            
        df_orders, df_items, df_unmatched_inv, invoice_row_matches = reconcile_data(sapo_path, supplier_path, invoice_path)
        
        output = create_4sheet_export_workbook(df_orders, df_items, df_unmatched_inv, invoice_row_matches, invoice_path)
        
        download_filename = f"Ket_qua_doi_chieu_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        return send_file(
            output,
            as_attachment=True,
            download_name=download_filename,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        
    except Exception as e:
        import traceback
        traceback.print_exc()
        return f"Đã xảy ra lỗi khi xuất file Excel: {str(e)}", 500

if __name__ == '__main__':
    # Listen on all interfaces so employees in LAN can access via http://<IP>:5000
    app.run(host='0.0.0.0', port=5000, debug=True)
